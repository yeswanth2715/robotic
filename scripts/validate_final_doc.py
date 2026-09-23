from pathlib import Path
from docx import Document

root=Path(__file__).resolve().parents[1]
doc=Document(root/'outputs'/'project5-final-dissertation.docx')
text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
checks=['4.1 Findings','4.2 Analysis','4.3 Discussion','WireSeg-36K','91 held-out','70.8 inference FPS','autonomous robot']
print('paragraphs=',len(doc.paragraphs),'tables=',len(doc.tables),'words=',len(text.split()))
for check in checks: print(check, check in text)
print('inline_shapes=',len(doc.inline_shapes))
