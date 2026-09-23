from pathlib import Path
from docx import Document

path=Path(__file__).resolve().parents[1]/'outputs'/'project5-final-dissertation.docx'
doc=Document(path)
text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
required=['4.1 Findings','4.2 Analysis','4.3 Discussion','WireSeg-36K','0.115','0.079','70.5 FPS']
print('docx_exists=',path.exists(),'tables=',len(doc.tables),'figures=',len(doc.inline_shapes),'words=',len(text.split()))
for value in required: print(value, value in text)
print('render_pdf=',(path.parent.parent/'qa'/'project5-final.pdf').exists())
